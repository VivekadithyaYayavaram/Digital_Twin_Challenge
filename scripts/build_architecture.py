"""Build submission/architecture_diagram.pptx (editable) and .pdf.

Run: python scripts/build_architecture.py   (or `make submission`)
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_helpers import (AMBER, AMBER_LIGHT, CARD, EHR, EHR_LIGHT, INK, LINE, MUTED, NAVY, SEN,  # noqa: E402
                          SEN_LIGHT, TEAL, TEAL_LIGHT, WHITE, arrow, blank, box, label_box, new_deck,
                          text, to_pdf)

OUT = Path(__file__).resolve().parents[1] / "submission"

M = 0.45            # left margin
GAP = 0.36          # horizontal gap (holds arrows)
W = (13.333 - 2 * M - 5 * GAP) / 6
X = [M + i * (W + GAP) for i in range(6)]


def lane(s, y, h, title, sub, fill):
    box(s, M - 0.15, y, 13.333 - 2 * M + 0.3, h, fill=fill, radius=0.12)
    text(s, M, y + 0.08, 9, 0.3, [[(title, {"bold": True, "size": 13, "color": NAVY}),
                                   ("   " + sub, {"size": 11, "color": MUTED})]])


def build() -> Path:
    prs = new_deck()
    s = blank(prs)
    text(s, M, 0.3, 12, 0.5, [[("SwasthTwin — system architecture", {"bold": True, "size": 26, "color": NAVY})]])
    text(s, M, 0.82, 12.4, 0.35, "Offline path builds and evaluates the model; runtime path keeps a per-patient "
         "Digital Twin and refreshes its 120-minute risk forecast as observations arrive.", size=12, color=MUTED)

    # ---------------- lane A: offline ---------------------------------------
    ya, ha = 1.32, 2.55
    lane(s, ya, ha, "OFFLINE · TRAINING-TIME", "run by `make data → make train → make evaluate`", CARD)
    yb = ya + 0.48; bh = ha - 0.62; half = (bh - 0.14) / 2
    label_box(s, X[0], yb, W, half, "Synthetic EHR", "80 fictional patients; HbA1c, labs, meds", EHR_LIGHT, EHR, line=EHR)
    label_box(s, X[0], yb + half + 0.14, W, half, "Synthetic sensors", "5-min CGM, HR, HRV, steps; sleep; meals", SEN_LIGHT, SEN, line=SEN)
    label_box(s, X[1], yb, W, bh, "Validate & align", "Schema, units, ranges, IDs linked; regular 5-min "
              "grid; impossible values → missing", WHITE, NAVY, line=LINE)
    label_box(s, X[2], yb, W, bh, "Features + labels", "Trailing windows only (≤ t). Label = any future "
              "CGM ≥180 mg/dL in (t, t+120 min]", WHITE, NAVY, line=LINE)
    label_box(s, X[3], yb, W, bh, "Patient-level split", "Train 60% / val 20% / test 20% of PATIENTS. "
              "Plus temporal hold-out (days 10+)", WHITE, NAVY, line=LINE)
    label_box(s, X[4], yb, W, bh, "Train & select", "Rule baseline, logistic reg., random forest, "
              "gradient boosting → pick by val PR-AUC; set threshold on val", TEAL_LIGHT, TEAL, line=TEAL)
    label_box(s, X[5], yb, W, half, "Evaluation & reports", "Test metrics, bootstrap CIs, "
              "ablation → reports/", WHITE, NAVY, line=LINE)
    label_box(s, X[5], yb + half + 0.14, W, half, "Model artifacts", "model_pipeline.joblib + "
              "metadata", WHITE, NAVY, line=NAVY)
    cy = yb + bh / 2
    for i in range(4):
        arrow(s, X[i] + W + 0.03, cy, X[i + 1] - 0.03, cy)
    arrow(s, X[4] + W + 0.03, cy - 0.25, X[5] - 0.03, yb + half / 2)
    arrow(s, X[4] + W + 0.03, cy + 0.25, X[5] - 0.03, yb + half + 0.14 + half / 2)

    # ---------------- lane B: runtime -----------------------------------------
    yr, hr = 4.08, 2.88
    lane(s, yr, hr, "RUNTIME · DIGITAL TWIN INFERENCE", "Streamlit dashboard, `make app`", TEAL_LIGHT)
    yb2 = yr + 0.48; bh2 = hr - 1.0; half2 = (bh2 - 0.14) / 2
    label_box(s, X[0], yb2, W, half2, "Patient EHR profile", "Historical, static context", EHR_LIGHT, EHR, line=EHR)
    label_box(s, X[0], yb2 + half2 + 0.14, W, half2, "Sensor feed (replay)", "Synthetic stream replay or manual "
              "entry — no live device yet", SEN_LIGHT, SEN, line=SEN, dashed=True)
    label_box(s, X[1], yb2, W, bh2, "Validate & align", "Same 5-min grid + range checks; staleness and "
              "coverage flags", WHITE, NAVY, line=LINE)
    label_box(s, X[2], yb2, W, bh2, "Rolling features", "SAME function as training → no train/serve skew",
              WHITE, NAVY, line=LINE)
    label_box(s, X[3], yb2, W, bh2, "Patient Digital Twin state", "EHR + observations ≤ clock + features + "
              "latest forecast + forecast history", NAVY, WHITE, line=NAVY, body_size=10.5)
    # make body text white on the navy box
    for shp in list(s.shapes)[-1:]:
        for p in shp.text_frame.paragraphs[1:]:
            for r in p.runs:
                r.font.color.rgb = WHITE
    label_box(s, X[4], yb2, W, bh2, "Trained pipeline", "Loaded artifact scores P(event in next 120 min)",
              TEAL_LIGHT, TEAL, line=TEAL)
    label_box(s, X[5], yb2, W, half2, "Risk + explanation", "Probability, prototype band, top inputs, "
              "data-quality warnings", WHITE, NAVY, line=NAVY)
    label_box(s, X[5], yb2 + half2 + 0.14, W, half2, "Clinician dashboard", "Timeline, fusion view, "
              "what-if, evaluation", WHITE, NAVY, line=NAVY)
    cy2 = yb2 + bh2 / 2
    arrow(s, X[0] + W + 0.03, yb2 + half2 / 2, X[1] - 0.03, cy2 - 0.25)
    arrow(s, X[0] + W + 0.03, yb2 + half2 + 0.14 + half2 / 2, X[1] - 0.03, cy2 + 0.25)
    for i in (1, 2, 3):
        arrow(s, X[i] + W + 0.03, cy2, X[i + 1] - 0.03, cy2)
    arrow(s, X[4] + W + 0.03, cy2, X[5] - 0.03, yb2 + half2 / 2)
    arrow(s, X[5] + W / 2, yb2 + half2 + 0.01, X[5] + W / 2, yb2 + half2 + 0.13)

    # artifact -> runtime model (crosses lanes)
    arrow(s, X[5] + 0.25, yb + bh + 0.02, X[4] + W - 0.25, yb2 - 0.03, color=TEAL, width=2.25)
    text(s, X[5] + 0.3, ya + ha + 0.0, 1.5, 0.22, "load saved model", size=10, color=TEAL, bold=True)

    # feedback loop: dashboard -> twin refresh
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Inches, Pt
    loop_y = yb2 + bh2 + 0.2
    xd, xt = X[5] + W / 2, X[3] + W / 2
    for (a, b, c, d) in ((xd, yb2 + bh2 + 0.01, xd, loop_y), (xd, loop_y, xt, loop_y)):
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(a), Inches(b), Inches(c), Inches(d))
        ln.line.color.rgb = AMBER; ln.line.width = Pt(2.0); ln.line.dash_style = 7
    arrow(s, xt, loop_y, xt, yb2 + bh2 + 0.02, color=AMBER, width=2.0, dashed=True)
    text(s, X[3] + W / 2 + 0.15, loop_y + 0.03, 6.0, 0.25,
         "advance clock · manual reading · what-if  →  recompute features  →  new forecast",
         size=10, color=AMBER, bold=True)

    # legend
    ly = 7.08
    box(s, M, ly + 0.03, 0.22, 0.16, fill=EHR_LIGHT, line=EHR)
    text(s, M + 0.3, ly, 1.8, 0.25, "historical EHR stream", size=10, color=MUTED)
    box(s, M + 2.1, ly + 0.03, 0.22, 0.16, fill=SEN_LIGHT, line=SEN)
    text(s, M + 2.4, ly, 1.8, 0.25, "dynamic sensor stream", size=10, color=MUTED)
    box(s, M + 4.2, ly + 0.03, 0.22, 0.16, fill=WHITE, line=MUTED, dashed=True)
    text(s, M + 4.5, ly, 3.4, 0.25, "dashed = simulated, not a live integration", size=10, color=MUTED)
    text(s, M + 7.9, ly, 4.6, 0.25, "All data synthetic · research prototype · not a medical device",
         size=10, color=AMBER, bold=True)

    s.notes_slide.notes_text_frame.text = (
        "Top lane runs offline: generate linked synthetic EHR and sensor data, validate, build leakage-safe "
        "features and future-window labels, split by patient, train candidates and pick by validation PR-AUC, "
        "save the pipeline and evaluate on held-out patients. Bottom lane is the live twin: the same feature "
        "function rebuilds the patient's state from observations up to the twin clock, the saved pipeline "
        "scores it, and the dashboard shows risk, explanation and warnings. Advancing the clock, entering a "
        "reading or running a what-if feeds back into the twin. No hospital system or wearable is integrated.")

    OUT.mkdir(parents=True, exist_ok=True)
    p = OUT / "architecture_diagram.pptx"
    prs.save(p)
    pdf = to_pdf(p)
    print("saved", p, pdf)
    return p


if __name__ == "__main__":
    build()
