"""Automated submission audit. Exit code 1 if any ERROR.

Checks: expected files exist; every relative Markdown link resolves; key metrics
quoted in README, model card and slides match reports/metrics.json; leftover
placeholders (reported as PENDING, not errors); no obvious secrets.

Run: python scripts/check_submission.py
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
errors: list[str] = []
pending: list[str] = []
ok: list[str] = []

EXPECTED = [
    "README.md", "LICENSE", "requirements.txt", "Makefile", ".gitignore", "app.py", "pyproject.toml",
    "config/model_config.yaml", "config/team_info.yaml",
    "data/ehr_patients.csv", "data/sensor_stream.csv.gz", "data/sleep_log.csv", "data/meal_log.csv",
    "artifacts/model_pipeline.joblib", "artifacts/model_metadata.json", "artifacts/patient_split.csv",
    "reports/metrics.json", "reports/metrics.md",
    "reports/figures/roc_pr_test.png", "reports/figures/calibration_test.png",
    "reports/figures/ablation_test.png", "reports/figures/permutation_importance_test.png",
    "reports/figures/confusion_test.png",
    "submission/architecture_diagram.pdf", "submission/architecture_diagram.pptx",
    "submission/Innovators_IIT_Kharagpur.pdf", "submission/Innovators_IIT_Kharagpur.pptx",
    "demo/DEMO_SCRIPT.md", "docs/model_card.md", "docs/data_provenance.md", "docs/data_dictionary.md",
    "docs/responsible_ai.md", "docs/architecture.md", "SUBMISSION_READINESS.md",
]


def check_files():
    missing = [f for f in EXPECTED if not (ROOT / f).exists()]
    (errors if missing else ok).append(f"expected files: {'missing ' + str(missing) if missing else 'all present'}")


def check_links():
    bad = []
    for md in ROOT.rglob("*.md"):
        if any(p in md.parts for p in (".venv", "node_modules", ".pytest_cache")):
            continue
        for target in re.findall(r"\]\(([^)\s]+)\)", md.read_text(encoding="utf-8")):
            if target.startswith(("http://", "https://", "mailto:", "#")):
                continue
            path = target.split("#")[0]
            if path and not (md.parent / path).resolve().exists():
                bad.append(f"{md.relative_to(ROOT)} -> {target}")
    (errors if bad else ok).append(f"relative links: {len(bad)} broken {bad[:10]}" if bad else "relative links: all resolve")


def check_metrics():
    m = json.loads((ROOT / "reports/metrics.json").read_text())
    best = m["selected_model"]
    t, r = m["test"][best], m["test"]["rule_baseline"]
    ci = m["bootstrap_ci_test"]["models"][best]["pr_auc"]
    must = {
        "selected PR-AUC": f"{t['pr_auc']:.3f}", "selected ROC-AUC": f"{t['roc_auc']:.3f}",
        "rule PR-AUC": f"{r['pr_auc']:.3f}", "selected PR CI": f"{ci[0]:.3f}–{ci[1]:.3f}",
        "onset PR-AUC": f"{m['test_onset_subset'][best]['pr_auc']:.3f}",
        "temporal PR-AUC": f"{m['temporal_holdout'][best]['pr_auc']:.3f}",
        "threshold": f"{m['operating_point']['threshold']:.3f}",
        "test sensitivity": f"{t['sensitivity']:.3f}",
    }
    for doc in ("README.md", "docs/model_card.md"):
        txt = (ROOT / doc).read_text(encoding="utf-8")
        miss = [k for k, v in must.items() if v not in txt]
        (errors if miss else ok).append(f"{doc} metrics match metrics.json" if not miss
                                        else f"{doc} missing/stale metrics: {miss}")
    try:
        from pptx import Presentation
        txt = " ".join(sh.text_frame.text for s in Presentation(ROOT / "submission/Innovators_IIT_Kharagpur.pptx").slides
                       for sh in s.shapes if sh.has_text_frame)
        need = [must["selected PR-AUC"], must["rule PR-AUC"], must["onset PR-AUC"], must["selected PR CI"]]
        miss = [v for v in need if v not in txt]
        (errors if miss else ok).append("slides metrics match metrics.json" if not miss
                                        else f"slides missing/stale metrics: {miss}")
        if Presentation(ROOT / "submission/Innovators_IIT_Kharagpur.pptx").slides.__len__() != 10:
            errors.append("presentation is not 10 slides")
    except ImportError:
        pending.append("python-pptx not installed: slide metric check skipped")


def check_placeholders():
    team = (ROOT / "config/team_info.yaml").read_text(encoding="utf-8")
    n = len(re.findall(r'"TODO', team))
    if n:
        pending.append(f"config/team_info.yaml has {n} TODO placeholders (team details)")
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    if "PENDING" in readme or "<VIDEO_URL>" in readme:
        pending.append("README: demo video URL still PENDING")
    if "<PUBLIC_REPO_URL>" in readme:
        pending.append("README: public repository URL placeholder")


def check_secrets():
    pat = re.compile(r"(sk-[A-Za-z0-9]{20,}|AKIA[0-9A-Z]{16}|-----BEGIN (RSA|OPENSSH) PRIVATE KEY-----|"
                     r"api[_-]?key\s*[:=]\s*['\"][A-Za-z0-9]{16,})", re.I)
    hits = []
    for f in ROOT.rglob("*"):
        if f.is_file() and f.suffix in {".py", ".md", ".yaml", ".toml", ".txt", ".json", ".cfg"} \
                and ".venv" not in f.parts:
            if pat.search(f.read_text(encoding="utf-8", errors="ignore")):
                hits.append(str(f.relative_to(ROOT)))
    (errors if hits else ok).append(f"secrets scan: {hits}" if hits else "secrets scan: none found")


def main() -> int:
    for fn in (check_files, check_links, check_metrics, check_placeholders, check_secrets):
        try:
            fn()
        except Exception as exc:  # report, keep going
            errors.append(f"{fn.__name__} crashed: {exc!r}")
    for x in ok:
        print("OK      ", x)
    for x in pending:
        print("PENDING ", x)
    for x in errors:
        print("ERROR   ", x)
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
