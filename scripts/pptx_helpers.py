"""Small python-pptx helpers shared by the diagram and presentation builders."""
from __future__ import annotations

import subprocess
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
NAVY2 = RGBColor(0x14, 0x2D, 0x52)
TEAL = RGBColor(0x0F, 0x9D, 0x8F)
TEAL_LIGHT = RGBColor(0xE3, 0xF4, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x0B, 0x1F, 0x3A)
MUTED = RGBColor(0x5B, 0x6B, 0x82)
LINE = RGBColor(0xC9, 0xD1, 0xDC)
CARD = RGBColor(0xF4, 0xF7, 0xFA)
EHR = RGBColor(0x5E, 0x46, 0x91)
EHR_LIGHT = RGBColor(0xEF, 0xE9, 0xF7)
SEN = RGBColor(0x24, 0x49, 0x7D)
SEN_LIGHT = RGBColor(0xE5, 0xEE, 0xF9)
AMBER = RGBColor(0x9A, 0x5B, 0x00)
AMBER_LIGHT = RGBColor(0xFC, 0xEF, 0xD9)
RED = RGBColor(0xA9, 0x32, 0x26)
FONT = "Calibri"


def new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def blank(prs, bg: RGBColor = WHITE):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def text(slide, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, name=None, line_spacing=1.05, margin=0.0):
    """runs: str, or list of paragraphs; each paragraph a str or list of (text, opts) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        items = para if isinstance(para, list) else [(para, {})]
        for t, o in items:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", FONT)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
        if isinstance(para, list) and para and para[0][1].get("space_after") is not None:
            p.space_after = Pt(para[0][1]["space_after"])
    return tb


def bullets(slide, x, y, w, h, items, size=15, color=INK, space=8):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.22))))
        pPr.set("indent", str(-Emu(Inches(0.22))))
        bu = etree.SubElement(pPr, qn("a:buChar"))
        bu.set("char", "•")
        parts = it if isinstance(it, list) else [(it, {})]
        for t, o in parts:
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", False)
            r.font.color.rgb = o.get("color", color)
    return tb


def box(slide, x, y, w, h, fill=CARD, line=None, radius=0.08, dashed=False, name=None, shadow=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.25)
        if dashed:
            shp.line.dash_style = 7  # dash
    if not shadow:
        sp = shp._element.spPr
        etree.SubElement(sp, qn("a:effectLst"))
    if name:
        shp.name = name
    shp.text_frame.text = ""
    return shp


def label_box(slide, x, y, w, h, title, body, fill, title_color, line=None, dashed=False,
              title_size=12, body_size=10.5, name=None):
    box(slide, x, y, w, h, fill=fill, line=line, dashed=dashed, name=name)
    text(slide, x + 0.1, y + 0.08, w - 0.2, h - 0.16,
         [[(title, {"bold": True, "size": title_size, "color": title_color})],
          [(body, {"size": body_size, "color": INK})]], line_spacing=1.0)


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.75, dashed=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dashed:
        c.line.dash_style = 7
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    return c


def circle_num(slide, x, y, d, label, fill=TEAL, color=WHITE, size=14):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = fill
    c.line.fill.background()
    etree.SubElement(c._element.spPr, qn("a:effectLst"))
    tf = c.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = color; r.font.name = FONT
    return c


def to_pdf(pptx: Path) -> Path:
    """Convert with LibreOffice (headless). Returns the PDF path."""
    out = pptx.parent
    cmd = ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out), str(pptx)]
    wrapper = Path("/mnt/skills/public/pptx/scripts/office/soffice.py")
    if wrapper.exists():  # sandbox-friendly wrapper when available
        cmd = ["python", str(wrapper), "--headless", "--convert-to", "pdf", "--outdir", str(out), str(pptx)]
    subprocess.run(cmd, check=True, capture_output=True, timeout=240)
    pdf = pptx.with_suffix(".pdf")
    if not pdf.exists():
        raise RuntimeError(f"PDF conversion failed for {pptx}")
    return pdf
